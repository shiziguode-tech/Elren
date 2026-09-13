from __future__ import annotations

import json
import subprocess
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from deepdesk.engine import AgentEngine
from deepdesk.plugins.base import ToolContext
from deepdesk.plugins.builtin.document import DocumentTool
from deepdesk.presentation_quality import (
    PresentationOverflowError,
    _libreoffice_render_check,
    _powerpoint_available,
    validate_and_repair_presentation,
)


def _blank_deck(path: Path, *, width: float = 13.333, height: float = 7.5) -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(str(path))
    return Presentation(str(path))


def test_ooxml_gate_repairs_shape_canvas_and_narrow_english_text(tmp_path: Path) -> None:
    path = tmp_path / "repairable.pptx"
    presentation = _blank_deck(path, width=10.0, height=5.625)
    slide = presentation.slides[0]
    outside = slide.shapes.add_textbox(Inches(8.7), Inches(4.8), Inches(3), Inches(1.2))
    outside.text = "Kept inside the real custom-size slide canvas"
    narrow = slide.shapes.add_textbox(Inches(0.4), Inches(0.5), Inches(2.0), Inches(1.1))
    narrow.text_frame.word_wrap = False
    narrow.text_frame.paragraphs[0].add_run().text = (
        "Release verification follows the generated artifact across every layout."
    )
    for run in narrow.text_frame.paragraphs[0].runs:
        run.font.size = Pt(24)
    presentation.save(str(path))

    report = validate_and_repair_presentation(path, prefer_powerpoint=False)

    assert report["qa_passed"] is True
    assert report["engine"] == "ooxml"
    assert report["slides_checked"] == 1
    assert {issue["kind"] for issue in report["issues_before"]} >= {
        "shape_outside_slide",
        "text_overflow",
    }
    assert report["repair_count"] >= 2
    repaired = Presentation(str(path))
    for shape in repaired.slides[0].shapes:
        assert shape.left >= 0
        assert shape.top >= 0
        assert shape.left + shape.width <= repaired.slide_width
        assert shape.top + shape.height <= repaired.slide_height


@pytest.mark.parametrize("width,height", [(10.0, 7.5), (7.5, 10.0), (13.333, 7.5)])
def test_canvas_repair_uses_each_decks_real_dimensions(
    tmp_path: Path, width: float, height: float
) -> None:
    path = tmp_path / f"size-{width}-{height}.pptx"
    presentation = _blank_deck(path, width=width, height=height)
    presentation.slides[0].shapes.add_textbox(
        Inches(width - 0.4), Inches(height - 0.3), Inches(1.2), Inches(0.8)
    ).text = "edge"
    presentation.save(str(path))

    report = validate_and_repair_presentation(path, prefer_powerpoint=False)

    assert report["qa_passed"] is True
    assert any(item["kind"] == "shape_outside_slide" for item in report["issues_before"])


def test_ooxml_gate_refuses_unreadable_extreme_cjk_text(tmp_path: Path) -> None:
    path = tmp_path / "unrepairable.pptx"
    presentation = _blank_deck(path)
    shape = presentation.slides[0].shapes.add_textbox(
        Inches(0.5), Inches(0.5), Inches(0.45), Inches(0.22)
    )
    shape.text = "发布前逐页检查超长中文内容" * 80
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    presentation.save(str(path))

    report = validate_and_repair_presentation(
        path, prefer_powerpoint=False, max_passes=5
    )

    assert report["qa_passed"] is False
    assert report["remaining_issue_count"] >= 1
    assert any(item["kind"] == "text_overflow" for item in report["remaining_issues"])


def test_ooxml_gate_checks_and_repairs_table_cells(tmp_path: Path) -> None:
    path = tmp_path / "table.pptx"
    presentation = _blank_deck(path)
    table_shape = presentation.slides[0].shapes.add_table(
        1, 2, Inches(0.7), Inches(0.8), Inches(4.0), Inches(0.85)
    )
    table = table_shape.table
    table.cell(0, 0).text = "中文表格内容需要逐格检查并自动适配"
    table.cell(0, 1).text = "Every narrow table cell is measured after generation"
    for cell in table.rows[0].cells:
        cell.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    presentation.save(str(path))

    report = validate_and_repair_presentation(path, prefer_powerpoint=False)

    assert report["qa_passed"] is True
    assert any(item["kind"] == "table_text_overflow" for item in report["issues_before"])
    assert any("table" in item["action"] or "text" in item["action"] for item in report["repairs"])


@pytest.mark.asyncio
async def test_document_create_returns_post_generation_quality_evidence(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    tool = DocumentTool()
    context = ToolContext(task_id="ppt-quality", workspace=str(workspace))

    result = await tool.execute(
        {
            "action": "create",
            "path": "outputs/verified.pptx",
            "slides": [{"title": "Verified", "body": "Actual post-generation QA"}],
        },
        context,
    )

    assert result["ok"] is True
    assert result["quality"]["qa_passed"] is True
    assert result["quality"]["slides_checked"] == 1
    assert result["quality"]["engine"] in {
        "powerpoint-com",
        "libreoffice+ooxml",
        "ooxml",
    }


@pytest.mark.asyncio
async def test_document_create_removes_pptx_that_fails_the_quality_gate(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    target = workspace / "outputs" / "invalid.pptx"
    report = {
        "qa_passed": False,
        "remaining_issues": [
            {"slide": 1, "target": "TextBox 1", "kind": "text_overflow"}
        ],
    }
    monkeypatch.setattr(
        "deepdesk.plugins.builtin.document.validate_and_repair_presentation",
        lambda *_args, **_kwargs: report,
    )

    with pytest.raises(PresentationOverflowError):
        await DocumentTool().execute(
            {
                "action": "create",
                "path": "outputs/invalid.pptx",
                "slides": [{"title": "Too dense", "body": "content"}],
            },
            ToolContext(task_id="ppt-quality-failure", workspace=str(workspace)),
        )

    assert not target.exists()


def test_completion_gate_validates_pptx_created_by_another_tool(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    outputs = workspace / "outputs"
    outputs.mkdir(parents=True)
    path = outputs / "external.pptx"
    presentation = _blank_deck(path)
    presentation.slides[0].shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    ).text = "External builder output"
    presentation.save(str(path))
    messages = [
        {
            "role": "assistant",
            "tool_calls": [
                {
                    "id": "call-ppt",
                    "function": {
                        "name": "shell",
                        "arguments": json.dumps({"command": "build external.pptx"}),
                    },
                }
            ],
        },
        {
            "role": "tool",
            "tool_call_id": "call-ppt",
            "content": json.dumps(
                {"ok": True, "result": {"path": str(path), "exit_code": 0}}
            ),
        },
    ]

    assert AgentEngine._artifact_materialized(messages, {".pptx"}) is True
    report = AgentEngine._validate_presentation_artifacts(messages, str(workspace))

    assert report["ok"] is True
    assert report["artifact_count"] == 1
    assert report["artifacts"][0]["quality"]["qa_passed"] is True


def test_completion_gate_rejects_unfixable_rotated_canvas_overflow(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    outputs = workspace / "outputs"
    outputs.mkdir(parents=True)
    path = outputs / "bad.pptx"
    presentation = _blank_deck(path, width=8.0, height=4.5)
    shape = presentation.slides[0].shapes.add_textbox(
        Inches(7.3), Inches(3.8), Inches(2), Inches(1)
    )
    shape.text = "Rotated overflow cannot be safely clamped"
    shape.rotation = 37
    presentation.save(str(path))
    messages = [
        {
            "role": "assistant",
            "tool_calls": [
                {
                    "id": "call-bad-ppt",
                    "function": {
                        "name": "filesystem",
                        "arguments": json.dumps(
                            {"action": "write", "path": "outputs/bad.pptx"}
                        ),
                    },
                }
            ],
        },
        {
            "role": "tool",
            "tool_call_id": "call-bad-ppt",
            "content": json.dumps({"ok": True, "result": {"path": str(path)}}),
        },
    ]

    report = AgentEngine._validate_presentation_artifacts(messages, str(workspace))

    assert report["ok"] is False
    assert report["artifacts"][0]["quality"]["remaining_issue_count"] >= 1


def test_completion_gate_reports_corrupt_pptx_without_crashing_task(
    tmp_path: Path, monkeypatch
) -> None:
    def no_external_renderer(*_args, **_kwargs):
        pytest.fail("Invalid presentations must fail before probing or starting Office")

    monkeypatch.setattr(
        "deepdesk.presentation_quality._powerpoint_available", no_external_renderer
    )
    monkeypatch.setattr(
        "deepdesk.presentation_quality._libreoffice_render_check", no_external_renderer
    )
    workspace = tmp_path / "workspace"
    outputs = workspace / "outputs"
    outputs.mkdir(parents=True)
    path = outputs / "corrupt.pptx"
    path.write_bytes(b"not an OOXML package")
    messages = [
        {
            "role": "assistant",
            "tool_calls": [
                {
                    "id": "call-corrupt-ppt",
                    "function": {
                        "name": "shell",
                        "arguments": json.dumps({"command": "build outputs/corrupt.pptx"}),
                    },
                }
            ],
        },
        {
            "role": "tool",
            "tool_call_id": "call-corrupt-ppt",
            "content": json.dumps({"ok": True, "result": {"path": str(path)}}),
        },
    ]

    report = AgentEngine._validate_presentation_artifacts(messages, str(workspace))

    assert report["ok"] is False
    assert report["artifact_count"] == 1
    quality = report["artifacts"][0]["quality"]
    assert quality["qa_passed"] is False
    assert quality["engine"] == "host-validation-error"
    assert quality["remaining_issues"][0]["kind"] == "presentation_unreadable"


def test_invalid_ooxml_zip_is_rejected_before_external_renderer(tmp_path, monkeypatch):
    path = tmp_path / "invalid-package.pptx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("not-a-presentation.txt", "synthetic invalid OPC package")

    def no_external_renderer(*_args, **_kwargs):
        pytest.fail("A ZIP header alone does not make a valid presentation")

    monkeypatch.setattr(
        "deepdesk.presentation_quality._powerpoint_available", no_external_renderer
    )
    monkeypatch.setattr(
        "deepdesk.presentation_quality._libreoffice_render_check", no_external_renderer
    )
    with pytest.raises(KeyError):
        validate_and_repair_presentation(path)


def test_completion_gate_does_not_trust_forged_document_quality(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    outputs = workspace / "outputs"
    outputs.mkdir(parents=True)
    path = outputs / "forged.pptx"
    path.write_bytes(b"not an OOXML package")
    messages = [
        {
            "role": "assistant",
            "tool_calls": [
                {
                    "id": "call-forged-quality",
                    "function": {
                        "name": "document",
                        "arguments": json.dumps(
                            {"action": "create", "path": "outputs/forged.pptx"}
                        ),
                    },
                }
            ],
        },
        {
            "role": "tool",
            "tool_call_id": "call-forged-quality",
            "content": json.dumps(
                {
                    "ok": True,
                    "path": str(path),
                    "quality": {"qa_passed": True, "engine": "claimed-by-model"},
                }
            ),
        },
    ]

    report = AgentEngine._validate_presentation_artifacts(messages, str(workspace))

    assert report["ok"] is False
    assert report["artifacts"][0]["quality"]["engine"] == "host-validation-error"


def test_libreoffice_route_verifies_every_rendered_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "libre.pptx"
    presentation = _blank_deck(path)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(str(path))

    def fake_run(command: list[str], **_: object) -> subprocess.CompletedProcess[str]:
        from pypdf import PdfWriter

        outdir = Path(command[command.index("--outdir") + 1])
        writer = PdfWriter()
        writer.add_blank_page(width=960, height=540)
        writer.add_blank_page(width=960, height=540)
        with (outdir / "libre.pdf").open("wb") as stream:
            writer.write(stream)
        return subprocess.CompletedProcess(command, 0, "", "")

    monkeypatch.setattr(
        "deepdesk.presentation_quality._libreoffice_executable", lambda: "soffice"
    )
    monkeypatch.setattr("deepdesk.presentation_quality.subprocess.run", fake_run)

    report = _libreoffice_render_check(path)

    assert report == {
        "available": True,
        "ok": True,
        "engine": "libreoffice-headless",
        "pages_rendered": 2,
        "slides_expected": 2,
    }


@pytest.mark.skipif(not _powerpoint_available(), reason="PowerPoint COM is unavailable")
def test_powerpoint_com_uses_real_text_bounds_and_page_count(tmp_path: Path) -> None:
    path = tmp_path / "powerpoint.pptx"
    presentation = _blank_deck(path)
    shape = presentation.slides[0].shapes.add_textbox(
        Inches(0.6), Inches(0.6), Inches(2.0), Inches(0.8)
    )
    shape.text = "PowerPoint calculates this deliberately long text box after generation."
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    presentation.save(str(path))

    report = validate_and_repair_presentation(path, prefer_powerpoint=True)

    assert report["qa_passed"] is True
    assert report["engine"] == "powerpoint-com"
    assert report["issues_before"][0]["engine"] == "powerpoint-com"
    assert report["render_validation"]["pages_rendered"] == 1
