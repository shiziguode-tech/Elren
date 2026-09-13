from pathlib import Path

import pytest

from deepdesk.plugins.base import ToolContext
from deepdesk.plugins.builtin.document import DocumentTool


@pytest.mark.asyncio
async def test_document_tool_creates_and_reads_all_supported_delivery_formats(tmp_path: Path):
    workspace = tmp_path / "workspace"
    attachment_root = tmp_path / "attachments"
    workspace.mkdir()
    attachment_root.mkdir()
    context = ToolContext(task_id="document-roundtrip", workspace=str(workspace))
    tool = DocumentTool(attachment_root)

    cases = {
        "outputs/report.pdf": {
            "title": "PDF report",
            "paragraphs": ["ELREN_PDF_MARKER", "Second paragraph"],
        },
        "outputs/report.docx": {
            "title": "Word report",
            "paragraphs": ["ELREN_DOCX_MARKER", "Second paragraph"],
        },
        "outputs/report.pptx": {
            "slides": [{"title": "Slide one", "body": "ELREN_PPTX_MARKER"}],
        },
        "outputs/report.xlsx": {
            "sheets": [{"name": "Results", "rows": [["Key", "Value"], ["marker", "ELREN_XLSX_MARKER"]]}],
        },
        "outputs/report.txt": {"text": "ELREN_TXT_MARKER\nPlain text result"},
    }

    for relative_path, payload in cases.items():
        created = await tool.execute(
            {"action": "create", "path": relative_path, **payload}, context
        )
        assert created["ok"] is True
        assert Path(created["path"]).is_file()
        inspected = await tool.execute(
            {"action": "inspect", "path": created["path"]}, context
        )
        marker = f"ELREN_{Path(relative_path).suffix[1:].upper()}_MARKER"
        assert marker in inspected["content"]
        assert inspected["format"] == Path(relative_path).suffix[1:]


@pytest.mark.asyncio
async def test_document_tool_reads_a_downloaded_attachment_but_cannot_write_outside_workspace(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    attachment_root = tmp_path / "attachments"
    workspace.mkdir()
    attachment_root.mkdir()
    attachment = attachment_root / "received.txt"
    attachment.write_text("FEISHU_ATTACHMENT_MARKER", encoding="utf-8")
    context = ToolContext(task_id="received-file", workspace=str(workspace))
    tool = DocumentTool(attachment_root)

    inspected = await tool.execute(
        {"action": "inspect", "path": str(attachment)}, context
    )
    assert "FEISHU_ATTACHMENT_MARKER" in inspected["content"]

    with pytest.raises(PermissionError):
        await tool.execute(
            {"action": "create", "path": str(tmp_path / "outside.docx"), "text": "blocked"},
            context,
        )


@pytest.mark.asyncio
async def test_document_tool_reads_legacy_rtf_doc_attachment(tmp_path: Path):
    workspace = tmp_path / "workspace"
    attachment_root = tmp_path / "attachments"
    workspace.mkdir()
    attachment_root.mkdir()
    attachment = attachment_root / "legacy.doc"
    attachment.write_bytes(
        b"{\\rtf1\\ansi\\deff0 {\\fonttbl {\\f0 Arial;}}\\f0 LEGACY_DOC_MARKER\\par Second line}"
    )
    tool = DocumentTool(attachment_root)
    context = ToolContext(task_id="legacy-doc", workspace=str(workspace))

    inspected = await tool.execute({"action": "inspect", "path": str(attachment)}, context)

    assert "LEGACY_DOC_MARKER" in inspected["content"]
    assert inspected["format"] == "doc"
    assert inspected["metadata"]["extractor"] == "striprtf"


@pytest.mark.asyncio
async def test_document_tool_reads_docx_container_with_doc_extension(tmp_path: Path):
    workspace = tmp_path / "workspace"
    attachment_root = tmp_path / "attachments"
    workspace.mkdir()
    attachment_root.mkdir()
    attachment = attachment_root / "misnamed.doc"
    from docx import Document

    document = Document()
    document.add_paragraph("MISNAMED_DOCX_MARKER")
    document.save(str(attachment))
    tool = DocumentTool(attachment_root)
    context = ToolContext(task_id="misnamed-docx", workspace=str(workspace))

    inspected = await tool.execute({"action": "inspect", "path": str(attachment)}, context)

    assert "MISNAMED_DOCX_MARKER" in inspected["content"]
    assert inspected["format"] == "doc"
    assert inspected["metadata"]["extractor"] == "python-docx"


@pytest.mark.asyncio
async def test_document_tool_reads_pptx_container_with_ppt_extension(tmp_path: Path):
    workspace = tmp_path / "workspace"
    attachment_root = tmp_path / "attachments"
    workspace.mkdir()
    attachment_root.mkdir()
    attachment = attachment_root / "legacy-name.ppt"
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Legacy extension"
    slide.placeholders[1].text = "PPT_LEGACY_CONTAINER_MARKER"
    presentation.save(str(attachment))
    tool = DocumentTool(attachment_root)
    context = ToolContext(task_id="misnamed-pptx", workspace=str(workspace))

    inspected = await tool.execute({"action": "inspect", "path": str(attachment)}, context)

    assert "PPT_LEGACY_CONTAINER_MARKER" in inspected["content"]
    assert inspected["format"] == "ppt"
    assert inspected["metadata"]["extractor"] == "python-pptx"


def test_document_tool_is_registered_for_the_agent():
    root = Path(__file__).resolve().parents[1]
    main = (root / "deepdesk" / "main.py").read_text(encoding="utf-8")
    engine = (root / "deepdesk" / "engine.py").read_text(encoding="utf-8")

    assert "DocumentTool(settings.screenshot_dir)" in main
    assert "document.inspect" in engine
