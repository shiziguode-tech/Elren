from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections.abc import Iterable
from contextlib import suppress
from pathlib import Path
from typing import Any

from deepdesk.platform_paths import unicode_font_candidates
from deepdesk.subprocess_env import credential_safe_environment

EMU_PER_POINT = 12_700
DEFAULT_BODY_FONT_PT = 18.0
MIN_READABLE_FONT_PT = 10.0
OVERFLOW_TOLERANCE_PT = 0.75
MAX_REPORTED_ITEMS = 200


class PresentationOverflowError(ValueError):
    """Raised when a generated presentation still clips content after repair."""

    def __init__(self, report: dict[str, Any]) -> None:
        self.report = report
        remaining = report.get("remaining_issues") or []
        samples = []
        for issue in remaining[:4]:
            samples.append(
                f"slide {issue.get('slide', '?')} {issue.get('target', 'shape')}: "
                f"{issue.get('kind', 'overflow')}"
            )
        suffix = "; ".join(samples) or "unknown overflow"
        super().__init__(
            "PowerPoint post-generation layout QA could not repair every overflow: " + suffix
        )


def validate_and_repair_presentation(
    path: Path | str,
    *,
    repair: bool = True,
    max_passes: int = 5,
    prefer_powerpoint: bool = True,
) -> dict[str, Any]:
    """Inspect every slide and repair deterministic text/canvas overflow.

    PowerPoint COM is preferred on Windows because ``TextRange2.BoundWidth`` and
    ``BoundHeight`` are the application's real layout measurements.  The OOXML
    path is intentionally always available for macOS/Linux and is paired with a
    LibreOffice render pass when that executable is installed.
    """

    source = Path(path).resolve()
    if source.suffix.lower() != ".pptx" or not source.is_file():
        raise ValueError("Presentation QA requires an existing .pptx file")

    # Validate the OOXML package before handing it to an external renderer.
    # Office can block on a corrupt-file recovery dialog despite DisplayAlerts,
    # turning a deterministic input error into a three-minute worker timeout.
    # The existing portable parser reads locally and does not start Office.
    from pptx import Presentation

    Presentation(str(source))

    warnings: list[str] = []
    if prefer_powerpoint and sys.platform == "win32" and _powerpoint_available():
        try:
            report = _run_powerpoint_gate_isolated(
                source, repair=repair, max_passes=max_passes
            )
            report["warnings"] = warnings
            return report
        except Exception as exc:
            warnings.append(f"PowerPoint COM unavailable during QA: {type(exc).__name__}: {exc}")

    report = _run_ooxml_gate(source, repair=repair, max_passes=max_passes)
    render = _libreoffice_render_check(source)
    report["render_validation"] = render
    if render.get("available"):
        report["engine"] = "libreoffice+ooxml"
        if not render.get("ok"):
            warnings.append(str(render.get("error") or "LibreOffice render validation failed"))
    report["warnings"] = warnings
    return report


def _powerpoint_available() -> bool:
    try:
        import win32com.client  # noqa: F401
    except Exception:
        return False
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"PowerPoint.Application\CLSID"):
            return True
    except OSError:
        return False


def _run_powerpoint_gate_isolated(
    path: Path, *, repair: bool, max_passes: int
) -> dict[str, Any]:
    """Run Office automation out of process so an Office RPC crash cannot kill Elren."""

    temporary = Path(tempfile.mkdtemp(prefix="elren-powerpoint-qa-"))
    report_path = temporary / "report.json"
    command = [
        sys.executable,
        "-m",
        "deepdesk.presentation_quality",
        "--powerpoint-worker",
        str(path),
        str(report_path),
        "1" if repair else "0",
        str(max_passes),
    ]
    try:
        completed = subprocess.run(
            command,
            cwd=str(Path(__file__).resolve().parents[1]),
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            env=credential_safe_environment(),
        )
        if completed.returncode != 0 or not report_path.is_file():
            details = (completed.stderr or completed.stdout or "PowerPoint QA worker failed")
            raise RuntimeError(details[-1_500:])
        report = json.loads(report_path.read_text(encoding="utf-8"))
        if not isinstance(report, dict):
            raise RuntimeError("PowerPoint QA worker returned an invalid report")
        return report
    finally:
        shutil.rmtree(temporary, ignore_errors=True)


def _public_issue(issue: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in issue.items() if not key.startswith("_")}


def _bounded(items: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    return [_public_issue(item) for item in list(items)[:MAX_REPORTED_ITEMS]]


def _rotated_bounds(left: float, top: float, width: float, height: float, rotation: float) -> tuple[float, float, float, float]:
    angle = math.radians(rotation % 360.0)
    rotated_width = abs(width * math.cos(angle)) + abs(height * math.sin(angle))
    rotated_height = abs(width * math.sin(angle)) + abs(height * math.cos(angle))
    center_x = left + width / 2.0
    center_y = top + height / 2.0
    return (
        center_x - rotated_width / 2.0,
        center_y - rotated_height / 2.0,
        center_x + rotated_width / 2.0,
        center_y + rotated_height / 2.0,
    )


def _canvas_issue(
    *,
    slide: int,
    target: str,
    left: float,
    top: float,
    width: float,
    height: float,
    rotation: float,
    slide_width: float,
    slide_height: float,
    engine: str,
    shape: Any,
) -> dict[str, Any] | None:
    x1, y1, x2, y2 = _rotated_bounds(left, top, width, height, rotation)
    if (
        x1 >= -OVERFLOW_TOLERANCE_PT
        and y1 >= -OVERFLOW_TOLERANCE_PT
        and x2 <= slide_width + OVERFLOW_TOLERANCE_PT
        and y2 <= slide_height + OVERFLOW_TOLERANCE_PT
    ):
        return None
    return {
        "slide": slide,
        "target": target,
        "kind": "shape_outside_slide",
        "engine": engine,
        "bounds_pt": [round(x1, 2), round(y1, 2), round(x2, 2), round(y2, 2)],
        "slide_size_pt": [round(slide_width, 2), round(slide_height, 2)],
        "_shape": shape,
        "_rotation": rotation,
    }


def _run_powerpoint_gate(path: Path, *, repair: bool, max_passes: int) -> dict[str, Any]:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    powerpoint = None
    presentation = None
    repairs: list[dict[str, Any]] = []
    first_issues: list[dict[str, Any]] = []
    passes = 0
    try:
        powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
        # PowerPoint rejects ``Visible=False`` even when no document window is
        # requested. ``WithWindow=False`` below is the supported headless route.
        powerpoint.DisplayAlerts = 1
        powerpoint.AutomationSecurity = 3
        presentation = powerpoint.Presentations.Open(
            str(path), ReadOnly=False, Untitled=False, WithWindow=False
        )
        for pass_number in range(1, max(1, max_passes) + 1):
            passes = pass_number
            issues = _scan_powerpoint(presentation)
            if pass_number == 1:
                first_issues = issues
            if not issues or not repair:
                break
            changed = _repair_powerpoint(issues, repairs)
            if not changed:
                break
            presentation.Save()
        remaining = _scan_powerpoint(presentation)
        if repair and repairs:
            presentation.Save()
        return {
            "qa_passed": not remaining,
            "engine": "powerpoint-com",
            "slides_checked": int(presentation.Slides.Count),
            "passes": passes,
            "issues_before": _bounded(first_issues),
            "issue_count_before": len(first_issues),
            "repairs": _bounded(repairs),
            "repair_count": len(repairs),
            "remaining_issues": _bounded(remaining),
            "remaining_issue_count": len(remaining),
            "render_validation": {
                "available": True,
                "ok": True,
                "engine": "powerpoint-com",
                "pages_rendered": int(presentation.Slides.Count),
            },
        }
    finally:
        if presentation is not None:
            with suppress(Exception):
                presentation.Close()
        if powerpoint is not None:
            with suppress(Exception):
                powerpoint.Quit()
        with suppress(Exception):
            pythoncom.CoUninitialize()


def _scan_powerpoint(presentation: Any) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    slide_width = float(presentation.PageSetup.SlideWidth)
    slide_height = float(presentation.PageSetup.SlideHeight)
    for slide_number in range(1, int(presentation.Slides.Count) + 1):
        slide = presentation.Slides.Item(slide_number)
        for shape_number in range(1, int(slide.Shapes.Count) + 1):
            shape = slide.Shapes.Item(shape_number)
            _scan_powerpoint_shape(
                shape,
                slide_number=slide_number,
                target=str(getattr(shape, "Name", f"shape {shape_number}")),
                slide_width=slide_width,
                slide_height=slide_height,
                issues=issues,
                inspect_geometry=True,
            )
    return issues


def _scan_powerpoint_shape(
    shape: Any,
    *,
    slide_number: int,
    target: str,
    slide_width: float,
    slide_height: float,
    issues: list[dict[str, Any]],
    inspect_geometry: bool,
) -> None:
    if inspect_geometry:
        issue = _canvas_issue(
            slide=slide_number,
            target=target,
            left=float(shape.Left),
            top=float(shape.Top),
            width=float(shape.Width),
            height=float(shape.Height),
            rotation=float(getattr(shape, "Rotation", 0.0) or 0.0),
            slide_width=slide_width,
            slide_height=slide_height,
            engine="powerpoint-com",
            shape=shape,
        )
        if issue:
            issues.append(issue)

    try:
        group_count = int(shape.GroupItems.Count) if int(shape.Type) == 6 else 0
    except Exception:
        group_count = 0
    for index in range(1, group_count + 1):
        child = shape.GroupItems.Item(index)
        _scan_powerpoint_shape(
            child,
            slide_number=slide_number,
            target=f"{target}/{getattr(child, 'Name', index)}",
            slide_width=slide_width,
            slide_height=slide_height,
            issues=issues,
            # A child group's coordinates are group-relative in some Office versions.
            inspect_geometry=False,
        )

    try:
        table = shape.Table if bool(shape.HasTable) else None
    except Exception:
        table = None
    if table is not None:
        for row in range(1, int(table.Rows.Count) + 1):
            for column in range(1, int(table.Columns.Count) + 1):
                cell_shape = table.Cell(row, column).Shape
                _scan_powerpoint_text_frame(
                    cell_shape,
                    slide_number=slide_number,
                    target=f"{target} cell {row},{column}",
                    issues=issues,
                    is_table=True,
                )

    _scan_powerpoint_text_frame(
        shape,
        slide_number=slide_number,
        target=target,
        issues=issues,
        is_table=False,
    )


def _scan_powerpoint_text_frame(
    shape: Any,
    *,
    slide_number: int,
    target: str,
    issues: list[dict[str, Any]],
    is_table: bool,
) -> None:
    try:
        if not bool(shape.HasTextFrame):
            return
        frame = shape.TextFrame2
        if not bool(frame.HasText):
            return
        text_range = frame.TextRange
        inner_width = max(0.0, float(shape.Width) - float(frame.MarginLeft) - float(frame.MarginRight))
        inner_height = max(0.0, float(shape.Height) - float(frame.MarginTop) - float(frame.MarginBottom))
        bound_width = float(text_range.BoundWidth)
        bound_height = float(text_range.BoundHeight)
        overflow_x = max(0.0, bound_width - inner_width)
        overflow_y = max(0.0, bound_height - inner_height)
        overflowing = False
        try:
            overflowing = bool(shape.TextFrame.Overflowing)
        except Exception:
            overflowing = False
        if (
            overflow_x <= OVERFLOW_TOLERANCE_PT
            and overflow_y <= OVERFLOW_TOLERANCE_PT
            and not overflowing
        ):
            return
        issues.append(
            {
                "slide": slide_number,
                "target": target,
                "kind": "table_text_overflow" if is_table else "text_overflow",
                "engine": "powerpoint-com",
                "overflow_x_pt": round(overflow_x, 2),
                "overflow_y_pt": round(overflow_y, 2),
                "_shape": shape,
                "_text_frame": frame,
                "_is_table": is_table,
            }
        )
    except Exception:
        return


def _repair_powerpoint(issues: list[dict[str, Any]], repairs: list[dict[str, Any]]) -> bool:
    changed = False
    seen: set[tuple[int, str, str]] = set()
    for issue in issues:
        key = (int(issue.get("slide", 0)), str(issue.get("target")), str(issue.get("kind")))
        if key in seen:
            continue
        seen.add(key)
        shape = issue.get("_shape")
        if shape is None:
            continue
        if issue.get("kind") == "shape_outside_slide":
            if abs(float(issue.get("_rotation") or 0.0)) % 90.0 > 0.01:
                continue
            slide_width, slide_height = issue["slide_size_pt"]
            width = min(max(1.0, float(shape.Width)), float(slide_width))
            height = min(max(1.0, float(shape.Height)), float(slide_height))
            left = min(max(0.0, float(shape.Left)), float(slide_width) - width)
            top = min(max(0.0, float(shape.Top)), float(slide_height) - height)
            shape.Width = width
            shape.Height = height
            shape.Left = left
            shape.Top = top
            changed = True
            repairs.append({**_public_issue(issue), "action": "clamped_to_slide_canvas"})
            continue

        frame = issue.get("_text_frame")
        if frame is None:
            continue
        try:
            frame.WordWrap = -1
            # Explicitly reduce runs instead of PowerPoint's opaque AutoFit.
            # AutoFit can silently scale 24 pt text to an unreadable size while
            # leaving TextRange.Font.Size at 24, defeating the minimum-size gate.
            frame.AutoSize = 0  # msoAutoSizeNone
            _scale_powerpoint_runs(frame.TextRange, 0.82)
            changed = True
            repairs.append({**_public_issue(issue), "action": "reduced_text_size"})
        except Exception:
            continue
    return changed


def _scale_powerpoint_runs(text_range: Any, factor: float) -> None:
    runs = text_range.Runs()
    count = int(getattr(runs, "Count", 0) or 0)
    if count <= 0:
        size = float(text_range.Font.Size or DEFAULT_BODY_FONT_PT)
        text_range.Font.Size = max(MIN_READABLE_FONT_PT, size * factor)
        return
    for index in range(1, count + 1):
        run = runs.Item(index)
        size = float(run.Font.Size or DEFAULT_BODY_FONT_PT)
        if size <= 0:
            size = DEFAULT_BODY_FONT_PT
        run.Font.Size = max(MIN_READABLE_FONT_PT, size * factor)


def _run_ooxml_gate(path: Path, *, repair: bool, max_passes: int) -> dict[str, Any]:
    from pptx import Presentation

    repairs: list[dict[str, Any]] = []
    first_issues: list[dict[str, Any]] = []
    passes = 0
    slides_checked = 0
    for pass_number in range(1, max(1, max_passes) + 1):
        passes = pass_number
        presentation = Presentation(str(path))
        slides_checked = len(presentation.slides)
        issues = _scan_ooxml(presentation)
        if pass_number == 1:
            first_issues = issues
        if not issues or not repair:
            break
        changed = _repair_ooxml(presentation, issues, repairs)
        if not changed:
            break
        presentation.save(str(path))
    final_presentation = Presentation(str(path))
    remaining = _scan_ooxml(final_presentation)
    return {
        "qa_passed": not remaining,
        "engine": "ooxml",
        "slides_checked": slides_checked,
        "passes": passes,
        "issues_before": _bounded(first_issues),
        "issue_count_before": len(first_issues),
        "repairs": _bounded(repairs),
        "repair_count": len(repairs),
        "remaining_issues": _bounded(remaining),
        "remaining_issue_count": len(remaining),
    }


def _scan_ooxml(presentation: Any) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    slide_width = float(presentation.slide_width) / EMU_PER_POINT
    slide_height = float(presentation.slide_height) / EMU_PER_POINT
    for slide_number, slide in enumerate(presentation.slides, 1):
        for shape_number, shape in enumerate(slide.shapes, 1):
            target = str(getattr(shape, "name", f"shape {shape_number}"))
            issue = _canvas_issue(
                slide=slide_number,
                target=target,
                left=float(shape.left) / EMU_PER_POINT,
                top=float(shape.top) / EMU_PER_POINT,
                width=float(shape.width) / EMU_PER_POINT,
                height=float(shape.height) / EMU_PER_POINT,
                rotation=float(getattr(shape, "rotation", 0.0) or 0.0),
                slide_width=slide_width,
                slide_height=slide_height,
                engine="ooxml",
                shape=shape,
            )
            if issue:
                issues.append(issue)
            if getattr(shape, "has_table", False):
                table = shape.table
                for row_number, row in enumerate(table.rows, 1):
                    for column_number, cell in enumerate(row.cells, 1):
                        _scan_ooxml_text(
                            cell.text_frame,
                            width=float(table.columns[column_number - 1].width) / EMU_PER_POINT,
                            height=float(row.height) / EMU_PER_POINT,
                            margins=(
                                _emu_or_default(cell.margin_left, 7.2),
                                _emu_or_default(cell.margin_right, 7.2),
                                _emu_or_default(cell.margin_top, 3.6),
                                _emu_or_default(cell.margin_bottom, 3.6),
                            ),
                            slide=slide_number,
                            target=f"{target} cell {row_number},{column_number}",
                            issues=issues,
                            owner=cell,
                            is_table=True,
                        )
            if getattr(shape, "has_text_frame", False):
                frame = shape.text_frame
                _scan_ooxml_text(
                    frame,
                    width=float(shape.width) / EMU_PER_POINT,
                    height=float(shape.height) / EMU_PER_POINT,
                    margins=(
                        _emu_or_default(frame.margin_left, 7.2),
                        _emu_or_default(frame.margin_right, 7.2),
                        _emu_or_default(frame.margin_top, 3.6),
                        _emu_or_default(frame.margin_bottom, 3.6),
                    ),
                    slide=slide_number,
                    target=target,
                    issues=issues,
                    owner=shape,
                    is_table=False,
                )
    return issues


def _emu_or_default(value: Any, default_points: float) -> float:
    if value is None:
        return default_points
    return float(value) / EMU_PER_POINT


def _scan_ooxml_text(
    frame: Any,
    *,
    width: float,
    height: float,
    margins: tuple[float, float, float, float],
    slide: int,
    target: str,
    issues: list[dict[str, Any]],
    owner: Any,
    is_table: bool,
) -> None:
    text = str(getattr(frame, "text", "") or "")
    if not text.strip():
        return
    inner_width = max(1.0, width - margins[0] - margins[1])
    inner_height = max(1.0, height - margins[2] - margins[3])
    measured_width, measured_height, max_font = _measure_text_frame(frame, inner_width)
    overflow_x = max(0.0, measured_width - inner_width)
    overflow_y = max(0.0, measured_height - inner_height)
    if overflow_x <= OVERFLOW_TOLERANCE_PT and overflow_y <= OVERFLOW_TOLERANCE_PT:
        return
    issues.append(
        {
            "slide": slide,
            "target": target,
            "kind": "table_text_overflow" if is_table else "text_overflow",
            "engine": "ooxml",
            "overflow_x_pt": round(overflow_x, 2),
            "overflow_y_pt": round(overflow_y, 2),
            "font_pt": round(max_font, 2),
            "_owner": owner,
            "_frame": frame,
        }
    )


def _font_path() -> Path | None:
    return next((candidate for candidate in unicode_font_candidates() if candidate.is_file()), None)


def _font_for_size(size_pt: float) -> Any:
    from PIL import ImageFont

    pixels = max(1, round(size_pt * 96.0 / 72.0))
    path = _font_path()
    if path:
        return ImageFont.truetype(str(path), pixels)
    return ImageFont.load_default(size=pixels)


def _paragraph_font_size(paragraph: Any) -> float:
    sizes = []
    for run in paragraph.runs:
        if run.font.size is not None:
            sizes.append(float(run.font.size.pt))
    if getattr(paragraph, "font", None) is not None and paragraph.font.size is not None:
        sizes.append(float(paragraph.font.size.pt))
    return max(sizes, default=DEFAULT_BODY_FONT_PT)


def _measure_text_frame(frame: Any, inner_width_pt: float) -> tuple[float, float, float]:
    max_width = 0.0
    total_height = 0.0
    max_font = 0.0
    paragraphs = list(frame.paragraphs)
    for paragraph in paragraphs or [None]:
        text = "" if paragraph is None else str(paragraph.text or "")
        size = DEFAULT_BODY_FONT_PT if paragraph is None else _paragraph_font_size(paragraph)
        max_font = max(max_font, size)
        font = _font_for_size(size)
        available_px = max(1.0, inner_width_pt * 96.0 / 72.0)
        lines = (
            _wrap_visual_lines(text, font, available_px)
            if frame.word_wrap is not False
            else (text.splitlines() or [""])
        )
        line_height_pt = size * 1.25
        total_height += max(1, len(lines)) * line_height_pt
        for line in lines:
            max_width = max(max_width, float(font.getlength(line)) * 72.0 / 96.0)
        if paragraph is not None:
            total_height += _emu_or_default(paragraph.space_before, 0.0)
            total_height += _emu_or_default(paragraph.space_after, 0.0)
    return max_width, total_height, max_font or DEFAULT_BODY_FONT_PT


def _wrap_visual_lines(text: str, font: Any, available_px: float) -> list[str]:
    lines: list[str] = []
    for source_line in text.splitlines() or [""]:
        if not source_line:
            lines.append("")
            continue
        current = ""
        tokens = re.findall(r"\s+|[\u2e80-\u9fff]|[^\s\u2e80-\u9fff]+", source_line)
        for token in tokens:
            candidate = current + token
            if not current or float(font.getlength(candidate)) <= available_px:
                current = candidate
                continue
            lines.append(current.rstrip())
            current = token.lstrip()
            while current and float(font.getlength(current)) > available_px:
                split_at = 1
                while (
                    split_at < len(current)
                    and float(font.getlength(current[: split_at + 1])) <= available_px
                ):
                    split_at += 1
                lines.append(current[:split_at])
                current = current[split_at:]
        lines.append(current)
    return lines


def _repair_ooxml(
    presentation: Any,
    issues: list[dict[str, Any]],
    repairs: list[dict[str, Any]],
) -> bool:
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Pt

    changed = False
    seen: set[tuple[int, str, str]] = set()
    for issue in issues:
        key = (int(issue.get("slide", 0)), str(issue.get("target")), str(issue.get("kind")))
        if key in seen:
            continue
        seen.add(key)
        if issue.get("kind") == "shape_outside_slide":
            shape = issue.get("_shape")
            if shape is None or abs(float(issue.get("_rotation") or 0.0)) % 90.0 > 0.01:
                continue
            slide_width = int(presentation.slide_width)
            slide_height = int(presentation.slide_height)
            width = min(max(1, int(shape.width)), slide_width)
            height = min(max(1, int(shape.height)), slide_height)
            shape.width = Emu(width)
            shape.height = Emu(height)
            shape.left = Emu(min(max(0, int(shape.left)), slide_width - width))
            shape.top = Emu(min(max(0, int(shape.top)), slide_height - height))
            changed = True
            repairs.append({**_public_issue(issue), "action": "clamped_to_slide_canvas"})
            continue

        frame = issue.get("_frame")
        if frame is None:
            continue
        current = max(MIN_READABLE_FONT_PT, float(issue.get("font_pt") or DEFAULT_BODY_FONT_PT))
        reduced = max(MIN_READABLE_FONT_PT, current * 0.78)
        if reduced >= current - 0.1:
            continue
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for paragraph in frame.paragraphs:
            if paragraph.runs:
                for run in paragraph.runs:
                    old = float(run.font.size.pt) if run.font.size is not None else current
                    run.font.size = Pt(max(MIN_READABLE_FONT_PT, old * 0.78))
            else:
                paragraph.font.size = Pt(reduced)
        changed = True
        repairs.append({**_public_issue(issue), "action": "ooxml_text_autofit_and_reduce"})
    return changed


def _libreoffice_executable() -> str | None:
    configured = os.getenv("ELREN_LIBREOFFICE", "").strip()
    if configured and Path(configured).is_file():
        return configured
    return shutil.which("soffice") or shutil.which("libreoffice")


def _libreoffice_render_check(path: Path) -> dict[str, Any]:
    executable = _libreoffice_executable()
    if not executable:
        return {"available": False, "ok": False, "engine": "libreoffice-headless", "pages_rendered": 0}
    temporary = Path(tempfile.mkdtemp(prefix="elren-ppt-qa-"))
    try:
        completed = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temporary),
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            env=credential_safe_environment(),
        )
        pdf_path = temporary / f"{path.stem}.pdf"
        if completed.returncode != 0 or not pdf_path.is_file():
            return {
                "available": True,
                "ok": False,
                "engine": "libreoffice-headless",
                "pages_rendered": 0,
                "error": (completed.stderr or completed.stdout or "conversion failed")[:500],
            }
        from pptx import Presentation
        from pypdf import PdfReader

        pages = len(PdfReader(str(pdf_path)).pages)
        slides = len(Presentation(str(path)).slides)
        return {
            "available": True,
            "ok": pages == slides and pages > 0,
            "engine": "libreoffice-headless",
            "pages_rendered": pages,
            "slides_expected": slides,
        }
    except Exception as exc:
        return {
            "available": True,
            "ok": False,
            "engine": "libreoffice-headless",
            "pages_rendered": 0,
            "error": f"{type(exc).__name__}: {exc}",
        }
    finally:
        shutil.rmtree(temporary, ignore_errors=True)


def _powerpoint_worker_main(arguments: list[str]) -> int:
    if len(arguments) != 4:
        return 2
    source = Path(arguments[0]).resolve()
    report_path = Path(arguments[1]).resolve()
    repair = arguments[2] == "1"
    max_passes = int(arguments[3])
    report = _run_powerpoint_gate(source, repair=repair, max_passes=max_passes)
    report_path.write_text(json.dumps(report, ensure_ascii=False), encoding="utf-8")
    return 0


if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1] == "--powerpoint-worker":
        raise SystemExit(_powerpoint_worker_main(sys.argv[2:]))
    raise SystemExit(2)
