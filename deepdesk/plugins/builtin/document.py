from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape
from zipfile import BadZipFile, ZipFile

from deepdesk.models import Risk
from deepdesk.platform_paths import unicode_font_candidates
from deepdesk.plugins.base import ToolContext, ToolPlugin, run_owned_thread
from deepdesk.presentation_quality import (
    PresentationOverflowError,
    validate_and_repair_presentation,
)
from deepdesk.subprocess_env import credential_safe_environment


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


class DocumentTool(ToolPlugin):
    name = "document"
    description = (
        "Inspect or create local PDF, Word (.doc/.docx), PowerPoint (.ppt/.pptx), Excel "
        "(.xlsx/.xls for inspection), and UTF-8 text files. Use inspect for uploaded "
        "attachments instead of trying to read binary Office files with filesystem. "
        "Create generated files under outputs/ and include the returned absolute path "
        "in the final answer when the file must be delivered to Feishu or Telegram."
    )
    parameters = {
        "type": "object",
        "properties": {
            "action": {"type": "string", "enum": ["inspect", "create"]},
            "path": {
                "type": "string",
                "description": "Absolute attachment path or workspace-relative file path",
            },
            "title": {"type": "string"},
            "text": {"type": "string"},
            "paragraphs": {"type": "array", "items": {"type": "string"}},
            "slides": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "body": {"type": "string"},
                    },
                    "additionalProperties": False,
                },
            },
            "sheets": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string"},
                        "rows": {
                            "type": "array",
                            "items": {"type": "array", "items": {}},
                        },
                    },
                    "required": ["rows"],
                    "additionalProperties": False,
                },
            },
        },
        "required": ["action", "path"],
        "additionalProperties": False,
    }

    inspect_extensions = {".pdf", ".doc", ".docx", ".ppt", ".pptx", ".xlsx", ".xls", ".txt"}
    create_extensions = {".pdf", ".docx", ".pptx", ".xlsx", ".txt"}

    def __init__(self, attachment_root: Path | None = None) -> None:
        self.attachment_root = attachment_root.resolve() if attachment_root else None

    def risk(self, arguments: dict[str, Any]) -> Risk:
        return Risk.SAFE if arguments.get("action") == "inspect" else Risk.MEDIUM

    def summarize(self, arguments: dict[str, Any]) -> str:
        action = "读取" if arguments.get("action") == "inspect" else "创建"
        return f"{action}文档：{arguments.get('path', '')}"

    async def execute(self, arguments: dict[str, Any], context: ToolContext) -> Any:
        return await run_owned_thread(self._execute_sync, arguments, context)

    @staticmethod
    def _within(path: Path, root: Path) -> bool:
        return path == root or root in path.parents

    def _resolve(self, requested: str, context: ToolContext, *, writing: bool) -> Path:
        from deepdesk.task_files import is_project_context, resolve_project_read

        if not writing and is_project_context(context):
            return resolve_project_read(requested, context)
        workspace = Path(context.workspace).resolve()
        raw = Path(str(requested or ""))
        path = raw.resolve() if raw.is_absolute() else (workspace / raw).resolve()
        if writing:
            if not self._within(path, workspace):
                raise PermissionError("Generated documents must stay inside the workspace")
        elif not self._within(path, workspace) and not (
            self.attachment_root and self._within(path, self.attachment_root)
        ):
            raise PermissionError("Document is outside the workspace and attachment directory")
        return path

    def _execute_sync(self, arguments: dict[str, Any], context: ToolContext) -> dict[str, Any]:
        action = str(arguments.get("action") or "")
        path = self._resolve(str(arguments.get("path") or ""), context, writing=action == "create")
        extension = path.suffix.lower()
        if action == "inspect":
            if extension not in self.inspect_extensions:
                raise ValueError(f"Unsupported document format: {extension or '[none]'}")
            if not path.is_file():
                raise FileNotFoundError(str(path))
            if path.stat().st_size > 50 * 1024 * 1024:
                raise ValueError("Document exceeds the 50 MB local inspection limit")
            return self._inspect(path)
        if action == "create":
            if extension not in self.create_extensions:
                raise ValueError(
                    "Create supports .pdf, .docx, .pptx, .xlsx and .txt; use .xlsx instead of legacy .xls"
                )
            path.parent.mkdir(parents=True, exist_ok=True)
            quality = self._create(path, arguments)
            result = {
                "ok": True,
                "path": str(path),
                "format": extension.lstrip("."),
                "bytes": path.stat().st_size,
            }
            if quality is not None:
                result["quality"] = quality
            return result
        raise ValueError(f"Unsupported document action: {action}")

    def _inspect(self, path: Path) -> dict[str, Any]:
        extension = path.suffix.lower()
        if extension == ".txt":
            text = path.read_text(encoding="utf-8", errors="replace")
            return self._result(path, text, {"characters": len(text)})
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = []
            for index, page in enumerate(reader.pages[:200], 1):
                pages.append(f"[Page {index}]\n{page.extract_text() or ''}")
            return self._result(path, "\n\n".join(pages), {"pages": len(reader.pages)})
        if extension == ".doc":
            return self._inspect_legacy_doc(path)
        if extension == ".docx":
            return self._inspect_docx(path)
        if extension == ".ppt":
            return self._inspect_legacy_ppt(path)
        if extension == ".pptx":
            return self._inspect_pptx(path)
        if extension == ".xlsx":
            self._validate_office_container(path)
            from openpyxl import load_workbook

            workbook = load_workbook(str(path), read_only=True, data_only=True)
            blocks = []
            sheet_meta = []
            cell_budget = 20_000
            for sheet in workbook.worksheets:
                blocks.append(f"[Sheet: {sheet.title}]")
                rows_read = 0
                for row in sheet.iter_rows(values_only=True):
                    if cell_budget <= 0:
                        break
                    values = ["" if value is None else str(value) for value in row]
                    blocks.append("\t".join(values))
                    cell_budget -= len(values)
                    rows_read += 1
                sheet_meta.append({"name": sheet.title, "rows_read": rows_read})
            workbook.close()
            return self._result(path, "\n".join(blocks), {"sheets": sheet_meta})
        if extension == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(str(path), on_demand=True)
            blocks = []
            sheet_meta = []
            cell_budget = 20_000
            for sheet in workbook.sheets():
                blocks.append(f"[Sheet: {sheet.name}]")
                rows_read = 0
                for row_index in range(sheet.nrows):
                    if cell_budget <= 0:
                        break
                    values = [str(sheet.cell_value(row_index, col)) for col in range(sheet.ncols)]
                    blocks.append("\t".join(values))
                    cell_budget -= len(values)
                    rows_read += 1
                sheet_meta.append({"name": sheet.name, "rows_read": rows_read})
            workbook.release_resources()
            return self._result(path, "\n".join(blocks), {"sheets": sheet_meta})
        raise ValueError(f"Unsupported document format: {extension}")

    def _inspect_docx(self, path: Path) -> dict[str, Any]:
        self._validate_office_container(path)
        from docx import Document

        document = Document(str(path))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        table_rows = 0
        for table_index, table in enumerate(document.tables, 1):
            blocks.append(f"[Table {table_index}]")
            for row in table.rows:
                blocks.append("\t".join(cell.text for cell in row.cells))
                table_rows += 1
        return self._result(
            path,
            "\n".join(blocks),
            {
                "extractor": "python-docx",
                "paragraphs": len(document.paragraphs),
                "tables": len(document.tables),
                "table_rows": table_rows,
            },
        )

    def _inspect_pptx(self, path: Path) -> dict[str, Any]:
        self._validate_office_container(path)
        from pptx import Presentation

        presentation = Presentation(str(path))
        blocks = []
        for slide_index, slide in enumerate(presentation.slides, 1):
            blocks.append(f"[Slide {slide_index}]")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    blocks.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        blocks.append("\t".join(cell.text for cell in row.cells))
        return self._result(
            path,
            "\n".join(blocks),
            {"slides": len(presentation.slides), "extractor": "python-pptx"},
        )

    def _inspect_legacy_ppt(self, path: Path) -> dict[str, Any]:
        data = path.read_bytes()
        if data.startswith(b"PK\x03\x04"):
            return self._inspect_pptx(path)
        if not data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise ValueError("The .ppt file is not a recognized PowerPoint presentation")
        errors: list[str] = []
        try:
            text, slides = self._inspect_legacy_ppt_with_powerpoint(path)
            return self._result(
                path, text, {"slides": slides, "extractor": "powerpoint-com-read-only"}
            )
        except Exception as exc:
            errors.append(f"PowerPoint: {type(exc).__name__}: {exc}")
        try:
            converted = self._convert_legacy_ppt_with_libreoffice(path)
            try:
                result = self._inspect_pptx(converted)
                result["path"] = str(path)
                result["format"] = "ppt"
                result["metadata"]["extractor"] = "libreoffice-to-python-pptx"
                return result
            finally:
                shutil.rmtree(converted.parent, ignore_errors=True)
        except Exception as exc:
            errors.append(f"LibreOffice: {type(exc).__name__}: {exc}")
        text = self._legacy_doc_printable_text(data)
        if text:
            return self._result(
                path,
                text,
                {"extractor": "binary-text-fallback", "conversion_errors": errors},
            )
        raise ValueError(
            "Legacy .ppt parsing requires Microsoft PowerPoint or LibreOffice; "
            + "; ".join(errors)
        )

    @staticmethod
    def _inspect_legacy_ppt_with_powerpoint(path: Path) -> tuple[str, int]:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        powerpoint = None
        presentation = None
        try:
            powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
            powerpoint.Visible = False
            powerpoint.DisplayAlerts = 1
            powerpoint.AutomationSecurity = 3
            presentation = powerpoint.Presentations.Open(
                str(path), ReadOnly=True, Untitled=False, WithWindow=False
            )
            blocks: list[str] = []
            for slide in presentation.Slides:
                blocks.append(f"[Slide {slide.SlideIndex}]")
                for shape in slide.Shapes:
                    try:
                        if shape.HasTextFrame and shape.TextFrame.HasText:
                            value = str(shape.TextFrame.TextRange.Text or "").strip()
                            if value:
                                blocks.append(value)
                        if shape.HasTable:
                            for row in range(1, shape.Table.Rows.Count + 1):
                                values = []
                                for column in range(1, shape.Table.Columns.Count + 1):
                                    values.append(
                                        str(
                                            shape.Table.Cell(row, column)
                                            .Shape.TextFrame.TextRange.Text
                                            or ""
                                        ).strip()
                                    )
                                blocks.append("\t".join(values))
                    except Exception:
                        continue
            return "\n".join(blocks), int(presentation.Slides.Count)
        finally:
            if presentation is not None:
                presentation.Close()
            if powerpoint is not None:
                powerpoint.Quit()
            pythoncom.CoUninitialize()

    @staticmethod
    def _convert_legacy_ppt_with_libreoffice(path: Path) -> Path:
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        if not executable:
            raise RuntimeError("LibreOffice is not installed")
        temporary = Path(tempfile.mkdtemp(prefix="elren-ppt-"))
        completed = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pptx",
                "--outdir",
                str(temporary),
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            env=credential_safe_environment(),
        )
        target = temporary / f"{path.stem}.pptx"
        if completed.returncode != 0 or not target.is_file():
            shutil.rmtree(temporary, ignore_errors=True)
            raise RuntimeError((completed.stderr or completed.stdout or "conversion failed")[:500])
        return target

    def _inspect_legacy_doc(self, path: Path) -> dict[str, Any]:
        """Extract legacy Word files without bringing Word to the foreground.

        Files named .doc can be OLE Word documents, RTF, HTML, or OOXML files
        with an old extension. Detecting the actual container keeps all four
        common cases usable.
        """
        data = path.read_bytes()
        if data.startswith(b"PK\x03\x04"):
            return self._inspect_docx(path)
        if data.lstrip().startswith(b"{\\rtf"):
            from striprtf.striprtf import rtf_to_text

            text = rtf_to_text(data.decode("latin-1", errors="replace"))
            return self._result(path, text, {"extractor": "striprtf"})

        header = data[:4096].lower()
        if b"<html" in header or b"<body" in header or b"<meta" in header:
            raw = data.decode("utf-8", errors="replace")
            extractor = _HTMLTextExtractor()
            extractor.feed(raw)
            text = "\n".join(extractor.parts)
            return self._result(path, text, {"extractor": "html"})

        if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            try:
                text = self._inspect_legacy_doc_with_word(path)
                return self._result(path, text, {"extractor": "word-com-read-only"})
            except Exception as exc:
                text = self._legacy_doc_printable_text(data)
                if text:
                    return self._result(
                        path,
                        text,
                        {
                            "extractor": "binary-text-fallback",
                            "word_error": f"{type(exc).__name__}: {exc}",
                        },
                    )
                raise ValueError(
                    "Legacy .doc parsing requires Microsoft Word on this computer, and no readable fallback text was found"
                ) from exc

        raise ValueError("The .doc file is not a recognized Word, RTF, HTML, or OOXML document")

    @staticmethod
    def _inspect_legacy_doc_with_word(path: Path) -> str:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        word = None
        document = None
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            # msoAutomationSecurityForceDisable: do not execute document macros.
            word.AutomationSecurity = 3
            document = word.Documents.Open(
                str(path),
                ConfirmConversions=False,
                ReadOnly=True,
                AddToRecentFiles=False,
                Visible=False,
                OpenAndRepair=False,
            )
            return str(document.Content.Text or "")
        finally:
            if document is not None:
                document.Close(False)
            if word is not None:
                word.Quit()
            pythoncom.CoUninitialize()

    @staticmethod
    def _legacy_doc_printable_text(data: bytes) -> str:
        candidates: list[str] = []
        for encoding in ("utf-16-le", "gb18030", "cp1252"):
            try:
                decoded = data.decode(encoding, errors="ignore")
            except LookupError:
                continue
            runs = re.findall(r"[\u0020-\u007e\u00a0-\u024f\u2e80-\u9fff]{6,}", decoded)
            value = "\n".join(run.strip() for run in runs if run.strip())
            if value:
                candidates.append(value)
        return max(candidates, key=len, default="")

    @staticmethod
    def _validate_office_container(path: Path) -> None:
        """Reject malformed or unusually expanded Office ZIP containers."""
        try:
            with ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10_000:
                    raise ValueError("Office document contains too many archive entries")
                if sum(entry.file_size for entry in entries) > 200 * 1024 * 1024:
                    raise ValueError("Office document expands beyond the 200 MB safety limit")
        except BadZipFile as exc:
            raise ValueError("Office document is not a valid OOXML container") from exc

    @staticmethod
    def _result(path: Path, content: str, metadata: dict[str, Any]) -> dict[str, Any]:
        truncated = len(content) > 100_000
        return {
            "ok": True,
            "path": str(path),
            "format": path.suffix.lower().lstrip("."),
            "content": content[:100_000] + ("\n[truncated]" if truncated else ""),
            "truncated": truncated,
            "metadata": metadata,
        }

    def _create(self, path: Path, arguments: dict[str, Any]) -> dict[str, Any] | None:
        extension = path.suffix.lower()
        title = str(arguments.get("title") or path.stem)
        paragraphs = [str(value) for value in arguments.get("paragraphs") or []]
        text = str(arguments.get("text") or "")
        if not paragraphs and text:
            paragraphs = text.splitlines() or [text]
        if extension == ".txt":
            path.write_text(text or "\n".join(paragraphs), encoding="utf-8")
            return None
        if extension == ".docx":
            from docx import Document

            document = Document()
            if title:
                document.add_heading(title, 0)
            for paragraph in paragraphs:
                document.add_paragraph(paragraph)
            document.save(str(path))
            return None
        if extension == ".pptx":
            from pptx import Presentation

            presentation = Presentation()
            slides = arguments.get("slides") or [{"title": title, "body": text or "\n".join(paragraphs)}]
            for item in slides:
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = str(item.get("title") or "")
                slide.placeholders[1].text = str(item.get("body") or "")
            presentation.save(str(path))
            quality = validate_and_repair_presentation(path, repair=True)
            if not quality.get("qa_passed"):
                path.unlink(missing_ok=True)
                raise PresentationOverflowError(quality)
            return quality
        if extension == ".xlsx":
            from openpyxl import Workbook

            workbook = Workbook()
            workbook.remove(workbook.active)
            sheets = arguments.get("sheets") or [
                {"name": "Sheet1", "rows": [[line] for line in (paragraphs or [text])]}
            ]
            for index, item in enumerate(sheets, 1):
                raw_name = str(item.get("name") or f"Sheet{index}")[:31]
                name = raw_name.translate({ord(char): "_" for char in "[]:*?/\\"}) or f"Sheet{index}"
                sheet = workbook.create_sheet(name)
                for row in item.get("rows") or []:
                    sheet.append(list(row))
            workbook.save(str(path))
            return None
        if extension == ".pdf":
            from reportlab.lib.enums import TA_LEFT
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

            font_name = "Helvetica"
            for candidate in unicode_font_candidates():
                if not candidate.is_file():
                    continue
                try:
                    font_name = "ElrenUnicode"
                    pdfmetrics.registerFont(TTFont(font_name, str(candidate)))
                    break
                except Exception:
                    font_name = "Helvetica"
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                "ElrenTitle", parent=styles["Title"], fontName=font_name, leading=24
            )
            body_style = ParagraphStyle(
                "ElrenBody", parent=styles["BodyText"], fontName=font_name, leading=16, alignment=TA_LEFT
            )
            document = SimpleDocTemplate(
                str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm
            )
            story = [Paragraph(escape(title), title_style), Spacer(1, 8 * mm)] if title else []
            for paragraph in paragraphs:
                story.extend([Paragraph(escape(paragraph).replace("\n", "<br/>"), body_style), Spacer(1, 3 * mm)])
            document.build(story)
            return None
